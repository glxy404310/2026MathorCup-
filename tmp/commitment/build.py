from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
import shutil
root=Path.cwd(); work=root/'tmp/commitment'; out=root/'solution_a/submission'
sigs=['exec-4d8bcfec-156c-46e9-8d02-ee2c1f21d525.png','exec-5a596c1b-1647-4b59-bff5-06e29940e8c6.png','exec-67eb4fcb-df56-4027-b2de-e074fd0fad04.png']
for i,f in enumerate(sigs,1): shutil.copy(Path('/Users/Admin1/.codex/generated_images/01a17bf3-1dae-7380-a0cf-baf3df9cab18')/f,work/f'clean{i}.png')
d=Document(next(root.glob('2026*/*承诺书*.docx')))
for idx,val in [(10,'A'),(11,'MCB2600730'),(12,'南京大学')]:
 p=d.paragraphs[idx]
 p.runs[1].text='  '+val
 for r in p.runs[2:]:r.text=''
for idx,i in [(14,1),(16,2),(18,3)]:
 p=d.paragraphs[idx]; p.clear(); p.paragraph_format.left_indent=Cm(6); p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1
 p.alignment=0
 p.add_run(f'队员{i}.  ')
 p.add_run().add_picture(str(work/f'clean{i}.png'),height=Cm(0.95))
p=d.paragraphs[22];p.clear();p.add_run('日期：2026 年 10 月 27 日');p.alignment=2
for p in d.paragraphs:
 for r in p.runs:
  r.font.name='Songti SC'
  r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Songti SC')
  if not r.font.size:r.font.size=Pt(12)
from docx.oxml import OxmlElement
for idx,p in enumerate(d.paragraphs):
 pf=p.paragraph_format
 pf.space_before=Pt(0);pf.space_after=Pt(0)
 pf.line_spacing=Pt(18 if idx not in [14,16,18] else 33)
 pr=p._p.get_or_add_pPr()
 sg=OxmlElement('w:snapToGrid');sg.set(qn('w:val'),'0');pr.append(sg)
 if idx in [14,16,18]:
  ind=pr.find(qn('w:ind'))
  if ind is not None:pr.remove(ind)
  pf.left_indent=Cm(6);pf.first_line_indent=Pt(0)
 if idx in [15,17,19,21]:pf.line_spacing=Pt(4)
d.save(out/'A2600730承诺书.docx')
