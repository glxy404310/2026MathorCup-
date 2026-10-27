from docx import Document
from copy import deepcopy
from pathlib import Path
import json
p=Path('solution_a/paper/啤酒风味建模提交稿.docx');d=Document(p)
ref='[13] OpenAI Codex，GPT-6.1 Sol，OpenAI，使用日期2026-10-23至2026-10-27。'
last=next(x for x in d.paragraphs if x.text.startswith('[12]'))
new=deepcopy(last._p)
from docx.text.paragraph import Paragraph
r=Paragraph(new,last._parent);r.text=ref;last._p.addnext(new)
u=next(x for x in d.paragraphs if x.text.startswith('本研究使用OpenAI Codex'))
u.text=u.text.replace('OpenAI Codex辅助','OpenAI Codex（GPT-6.1 Sol）[13]辅助')
d.save(p)
Path('solution_a/docs/use_statement.md').write_text('# 使用声明\n\n'+u.text+'\n')
p=Path('solution_a/paper/manuscript.json');m=json.loads(p.read_text());b=m['blocks'];i=max(i for i,x in enumerate(b) if x.get('type')=='reference');b.insert(i+1,{'type':'reference','text':ref})
for x in b:
 if x.get('text','').startswith('本研究使用OpenAI Codex'):x['text']=u.text
p.write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n')
