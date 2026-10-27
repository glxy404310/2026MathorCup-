from pathlib import Path
import re,json,shutil,zipfile,hashlib
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
ROOT=Path.cwd(); W=ROOT/'tmp/humanizer'; BACK=W/'before'; BACK.mkdir(exist_ok=True)
paths=[Path('solution_a/paper/啤酒风味建模提交稿.docx'),Path('solution_a/submission/A2600730.pdf'),Path('solution_a/submission/A2600730附件.zip'),Path('solution_a/submission/AI工具使用详情说明.docx'),Path('solution_a/submission/AI工具使用详情说明.pdf')]
for p in paths:
 target=BACK/p.name
 if not target.exists():shutil.copy2(p,target)
repls={
'针对啤酒化学测量存在缺失、成分相关性强及部分风格样本较少的问题，研究感官预测、关键化合物分析和风格迁移。':'本文研究啤酒感官预测、关键化合物分析和风格迁移，重点处理化学测量缺失、成分相关性强及部分风格样本较少的问题。',
'按啤酒编号核对附件后，采用250款啤酒、45项感官属性；':'按啤酒编号核对附件，选取250款啤酒的45项感官属性作为研究对象；',
'所选谱目标距离缩短中位数':'所选化学谱到目标风格的距离缩短幅度中位数',
'按beer进行一对一连接。':'按beer编号一对一关联各表。',
'因此按实物矩阵采用Spearman，而不沿用题面对Pearson的文字称谓。':'因此，本文依据附件矩阵采用Spearman相关，题面所述Pearson相关与该矩阵不符。',
'第一问检验真实啤酒的折外预测误差；':'第一问评价真实啤酒样本的折外预测误差；',
'比较目标接近度和保护约束之间的取舍。':'比较接近目标风格的程度与满足保护约束之间的取舍。',
'故能缓和外推':'因此能够减小极端外推的影响',
'这些弱维度不宜承担强感官保证。':'这些预测较弱的维度尚不足以支持对实际感官效果的可靠判断。',
'但不按同批折外结果逐维拼接赢家，以免抬高评价。':'但不依据同批折外结果逐维选择最优模型再组合，以免高估预测性能。',
'未据此调参或逐目标拼接赢家。':'未根据该结果调参，也未逐目标选择最优模型后组合。',
'用于检查模型家族及部署范围':'用于比较模型类型并考察适用范围',
'不能由酸度等强项推及所有感官':'不能将酸度等属性的预测效果推广至所有感官属性',
'故不将缺失指示概括为统一增益。':'因此，缺失指示并未在所有情景下改善预测。',
'网络结构、候选稳定性和化学气味证据分别报告，相关清单不作为因果效应或独立测试性能。':'网络结构、候选稳定性和化学气味证据分别报告；候选清单仅描述相关性，不能据此推断因果效应或独立测试性能。',
'不以后验文献解释或预计算GBM/SHAP改选。':'不根据事后查阅的文献或预计算GBM/SHAP结果调整候选。',
'该补充不改全样本相关前三；':'该补充分析不改变全样本相关排名的前三项；',
'候选显著依赖关联定义和风格构成。':'候选名单随关联定义和风格构成明显变化。',
'不算外部重复验证。':'不能作为外部重复验证。',
'铁—金属味也需限定证据口径：':'铁与金属味的联系也需结合试验条件解释：',
'阈值缺失处不补造，外部试验阈值不覆盖附件值。':'缺失阈值保持空缺，也不以外部试验阈值替换附件值。',
'表6-3给出代表物':'表6-3列出代表性化合物',
'均不识别因果效应。':'均不能用于识别因果效应。',
'三问的联系应限于证据核对和缺陷变量解释。':'三问之间仅通过证据核对和缺陷变量解释建立联系。',
'生成检验分布贴近程度和新组合，迁移则在化学与感官约束下比较可行候选，两者分别评价。':'生成部分评价分布接近程度及新组合，迁移部分则在化学与感官约束下比较可行候选。两部分分别评价。',
'小风格样本不足以稳定估计226维协方差。':'样本较少的风格不足以稳定估计226维协方差。',
'整列缺失时回退全局训练中位数。':'整列缺失时改用全局训练中位数。',
'两者都不能补足未观测尾部。':'两者都无法补充未观测到的尾部分布信息。',
'通过后按平均能量选择，外层评价完整选择流程。':'满足条件后按平均能量差异选择方法，再由外层验证评价整个选择流程。',
'迁移沿用239款训练酒拟合的高斯秩变换和Tripel几何，使各方案使用同一距离。':'迁移沿用239款训练酒拟合的高斯秩变换及Tripel的几何表示，使各方案的距离可比。',
'以每款源酒原已保存迁移谱的改善幅度Gᵢ锚为参照':'以每款源酒此前保存的迁移谱改善幅度Gᵢ锚为参照',
'故不能把化学通过概括为四项浓度全部下降。':'因此，通过化学约束检查并不意味着四项浓度均下降。',
'本轮另以相同参数、不同种子重拟合两种树作为选型对照':'另以相同参数、不同种子重新拟合两种树作为选型对照',
'故仍作备选诊断。':'因此，仅将其保留为备选诊断结果。',
'候选存在和最终通过分开报告；':'是否存在可行候选与是否通过最终复核分别报告；',
'小风格样本偏少':'部分风格的样本偏少',
'10/11是模型检查通过数':'10/11表示通过模型检查的比例',
'表B-1沿用绝对Spearman相关排名，不按文献证据后验改选。':'表B-1沿用绝对Spearman相关排名，不根据事后查阅的文献调整候选。',
'阈值只用于对应的校准浓度；空缺阈值不补造。':'阈值仅与对应的校准浓度比较；缺失阈值保持空缺。',
'两个Q95为预留当前模型及新种子旧协议模型的配对预测差分位数；':'两个Q95分别为预留当前模型、新种子旧协议模型的配对预测差分位数；',
'不删除失败的beer 96。':'同时保留未通过复核的beer 96。',
'首次完整重建在新解压目录执行：':'首次完整重建时，在新解压的目录中执行：',
'程序对非空历史或补充输出拒绝覆盖；重新试验应另建副本。':'程序不会覆盖非空的历史或补充输出目录；重新试验时应另建副本。',
}
P=paths[0];d=Document(P);changes=[]
for idx,p in enumerate(d.paragraphs):
 old=p.text;new=old
 if idx<7 or idx>=45:
  if p._p.xpath('.//m:oMath') or p._p.xpath('.//w:instrText'):continue
  for a,b in repls.items():new=new.replace(a,b)
 if new!=old:
  assert re.findall(r'\d+(?:\.\d+)?',old)==re.findall(r'\d+(?:\.\d+)?',new),(idx,old,new)
  ppr=p._p.pPr
  keep=[e for e in p._p if e.tag in [qn('w:bookmarkStart'),qn('w:bookmarkEnd')]]
  props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
  for e in list(p._p):
   if e is not ppr and e not in keep:p._p.remove(e)
  r=p.add_run(new)
  if props is not None:r._r.insert(0,props)
  changes.append({'index':idx,'before':old,'after':new})
# Move the existing AI disclosure after references, retaining its anchors and content.
ps=d.paragraphs;uh=next(p for p in ps[45:] if p.text=='使用声明');ut=ps[ps.index(uh)+1] if False else next(p for p in ps if p.text.startswith('本研究使用OpenAI Codex'))
lastref=next(p for p in ps if p.text.startswith('[12]'))
lastref._p.addnext(uh._p);uh._p.addnext(ut._p)
# Keep the existing contents entries in the same order as the body.
toc_u=next(p for p in ps[:45] if p.text.startswith('使用声明'));toc_r=next(p for p in ps[:45] if p.text.startswith('参考文献'));toc_r._p.addnext(toc_u._p)
d.save(P)
# Synchronize prose-bearing JSON sources; leave numbers, equations and tables untouched.
def walk(v):
 if isinstance(v,dict):
  return {k:(v2 if k in ['latex','rows','headers','widths'] else walk(v2)) for k,v2 in v.items()}
 if isinstance(v,list):return [walk(x) for x in v]
 if isinstance(v,str):
  for a,b in repls.items():v=v.replace(a,b)
 return v
json_changed=[]
for p in [Path('solution_a/paper/manuscript.json'),*Path('solution_a/paper/sections').glob('*.json'),Path('solution_a/paper/figure_manifest.json')]:
 orig=json.loads(p.read_text());new=walk(orig)
 if p.name=='manuscript.json':
  bs=new['blocks'];j=next(i for i,b in enumerate(bs) if b.get('text')=='使用声明');pair=bs[j:j+2];del bs[j:j+2];k=max(i for i,b in enumerate(bs) if b.get('type')=='reference');bs[k+1:k+1]=pair
 if orig!=new:
  dest=BACK/p.relative_to('solution_a');dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest)
  p.write_text(json.dumps(new,ensure_ascii=False,indent=2)+'\n');json_changed.append(str(p))
# Precise editorial corrections in support prose, preserving code, links and history.
support_repls={
'论文定稿采用口径见':'论文定稿采用的方案见',
'其指标只作诊断':'其指标仅用于诊断',
'不能直接塞进本轮模型':'不能直接输入本轮模型',
'不得未经依据将250样本裁剪为226样本':'没有依据时，不应将250个样本删减为226个样本',
'不能无检查拼接':'合并前须核对字段对应关系',
'而不是只展示最漂亮的几个目标':'不能仅展示结果较好的目标',
'不依据同一批折外结果逐维挑选赢家拼接主模型':'不依据同一批折外结果逐维挑选最优模型后组合为主模型',
'不能重新制造一份从未接触过的独立数据':'不能使这些数据重新成为未经使用的独立测试集',
'实际跑通':'完成运行',
'响应性不证明预测方向正确':'存在响应并不证明预测方向正确',
'不能把载入工件等同于通过部署资格':'能够载入模型文件不意味着该模型符合部署条件',
}
reviewed=[];support_changes=[]
for p in [Path('solution_a/README.md'),*sorted(Path('solution_a/docs').glob('*.md')),Path('solution_a/supplement/README.txt')]:
 old=p.read_text();ls=old.splitlines(keepends=True);new=[];fenced=False
 for line in ls:
  if line.lstrip().startswith('```'):fenced=not fenced;new.append(line);continue
  if fenced or line.lstrip().startswith(('#','|')):new.append(line);continue
  parts=re.split(r'(`[^`]*`|\[[^\]]*\]\([^)]*\)|https?://\S+)',line)
  for i in range(0,len(parts),2):
   for a,b in support_repls.items():parts[i]=parts[i].replace(a,b)
  new.append(''.join(parts))
 new=''.join(new);reviewed.append(str(p))
 if old!=new:
  assert re.findall(r'\d+(?:\.\d+)?',old)==re.findall(r'\d+(?:\.\d+)?',new)
  dest=BACK/p.relative_to('solution_a');dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest);p.write_text(new)
  support_changes.append(str(p))
report={'main_changes':changes,'support_scanned':reviewed,'support_changed':support_changes,'json_synchronized':json_changed,'protected':'Model programs, numerical data, equations, tables, citation identities and official commitment unchanged.'}
(W/'edit_log.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print('Main paragraphs edited:',len(changes),'support files scanned:',len(reviewed),'support files edited:',len(support_changes))
