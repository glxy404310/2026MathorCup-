from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
D=Document();s=D.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(2.3);s.right_margin=Cm(2.3)
style=D.styles['Normal'];style.font.name='Songti SC';style.font.size=Pt(11);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Songti SC');style.paragraph_format.line_spacing=Pt(16);style.paragraph_format.space_after=Pt(5)
p=D.add_paragraph();p.alignment=1;r=p.add_run('AI工具使用详情说明');r.bold=True;r.font.size=Pt(17)
def h(t):
 p=D.add_paragraph();p.add_run(t).bold=True

def p(t):D.add_paragraph(t)
h('一、工具及使用范围')
p('本参赛队使用 OpenAI Codex（模型：GPT-6.1 Sol；开发机构：OpenAI），主要辅助完成附件数据清洗、数据预处理、回归计算程序实现与计算结果核验。相关工作记录形成于2026年10月23日至27日。')
h('二、主要应用环节')
p('1. 数据清洗与核对：辅助检查样本编号、重复记录、缺失值、字段类型和数值范围，核对不同附件之间的关联关系及变量名称。处理过程中保留原始数据，按 beer 编号关联样本，并记录缺失和异常值位置，避免因条码精度或字段顺序差异造成错误合并。')
p('2. 数据预处理：辅助实现字段映射、缺失值填补、标准化及稳健处理程序，检查训练集与验证集的分离。插补、缩放等参数在训练折内拟合，再应用于验证折；不以AI生成数值替代原始观测。')
p('3. 回归计算与核验：辅助编写、调试岭回归和树模型等回归程序，执行交叉验证，汇总 MAE、RMSE、R² 等指标，并排查数组维度、字段对应及结果导出问题。预测结果来自程序对实际附件数据的运算，不直接采用工具对预测数值的文字推测。')
p('此外，Codex用于计算结果图表整理，以及论文文字和排版的辅助修订。')
h('三、关键任务记录摘要')
p('原始提示词与回复未单独保存。以下依据项目中的数据审计、阶段结果和程序记录整理，为任务摘要，不是原始对话的逐字记录。')
p('数据清洗任务：检查附件结构、关联键和缺失情况。记录结果包括250个样本的键值核对、化学数据73处缺失，以及部分附件条码精度问题；相关检查保存在数据审计记录和输出文件中。')
p('数据处理任务：核对字段映射和相关系数口径。记录结果包括感官字段对应关系及给定相关矩阵的数值复算，处理程序采用显式字段匹配，避免按位置直接拼接。')
p('回归任务：实现和检查回归训练、折外预测及误差汇总。记录结果包括均值基线、岭回归和树回归的比较，以及缺失压力和稳健处理检查；具体数值以保存的预测文件和运行结果为依据。')
h('四、采纳内容与核验处理')
p('采纳内容主要为数据处理和回归程序的辅助实现、错误排查及指标汇总。项目记录中形成的修正包括调整样本关联键、纠正字段对应、将预处理限制在训练折内，以及复核预测与汇总指标的一致性。上述程序核验不等同于人工复核；参赛队对最终采纳内容、论文表述及提交结果的正确性负责。')

D.save('solution_a/submission/AI工具使用详情说明.docx')
