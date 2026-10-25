"""Assemble the editable submission paper from verified section and result files."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from native_math import to_omml

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
REPO = ROOT.parent
REFERENCE = next((REPO / "2026年第七届MathorCup数学应用挑战赛—大数据竞赛初赛赛题").glob("*论文模板.docx"))
FINAL = HERE / "啤酒风味建模提交稿.docx"
(HERE/'qa').mkdir(parents=True,exist_ok=True)
from submission_content import add_submission_appendices
from supplement_sections import add_main_sections

core = json.loads((HERE / "sections/core.json").read_text(encoding="utf-8"))
figs = json.loads((HERE / "figure_manifest.json").read_text(encoding="utf-8"))["figures"]
FIGURES = {Path(x["relativepath"]).stem: x for x in figs}


def paragraph(text):
    return {"type": "paragraph", "text": text}


def heading(text, level=1):
    return {"type": "heading", "level": level, "text": text}


def table(caption, headers, rows, widths=None, compact=False):
    return {"type": "table", "caption": caption, "headers": headers,
            "rows": rows, "widths": widths, "compact": compact}


q1, q2, q3 = [json.loads((HERE / f"sections/q{i}.json").read_text(encoding="utf-8"))["blocks"] for i in [1, 2, 3]]
q1, q2, q3 = add_main_sections(q1, q2, q3, ROOT)
q1.insert(next(i for i, b in enumerate(q1) if b.get("text", "").startswith("5.6 ")), {"type": "figure", "key": "figure2_q1"})
# Place figures close to the results they support.
q2.insert(next(i for i, b in enumerate(q2) if b.get("text", "").startswith("6.3 ")), {"type": "figure", "key": "figure3_q2"})
q3.insert(next(i for i, b in enumerate(q3) if b.get("text", "").startswith("7.5 ")), {"type": "figure", "key": "figure4_generation"})
q3.insert(next(i for i, b in enumerate(q3) if b.get("text", "").startswith("7.5 ")), {"type": "figure", "key": "figure5_profiles"})
q3.insert(next(i for i, b in enumerate(q3) if b.get("text", "").startswith("7.10 ")), {"type": "figure", "key": "figure6_transfer"})
blocks = deepcopy(core["opening"] + q1 + q2 + q3 + core["ending"])

refs = json.loads((HERE / "chemical_references.json").read_text(encoding="utf-8"))["references"]
refs.extend([
    {"number": 2, "authors": ["Pierre Geurts", "Damien Ernst", "Louis Wehenkel"],
     "title": "Extremely randomized trees", "journal": "Machine Learning", "year": 2006,
     "volume": "63", "pages": "3–42", "doi": "10.1007/s10994-006-6226-1",
     "url": "https://doi.org/10.1007/s10994-006-6226-1"},
    {"number": 3, "authors": ["Arthur E. Hoerl", "Robert W. Kennard"],
     "title": "Ridge Regression: Biased Estimation for Nonorthogonal Problems", "journal": "Technometrics",
     "year": 1970, "volume": "12", "issue": "1", "pages": "55–67", "doi": "10.1080/00401706.1970.10488634",
     "url": "https://www.stat.cmu.edu/technometrics/70-79/VOL-12-01/v1201055.pdf"},
    {"number": 4, "authors": ["Han Liu", "John Lafferty", "Larry Wasserman"],
     "title": "The Nonparanormal: Semiparametric Estimation of High Dimensional Undirected Graphs",
     "journal": "Journal of Machine Learning Research", "year": 2009, "volume": "10", "pages": "2295–2328",
     "doi": None, "url": "https://www.jmlr.org/papers/v10/liu09a.html"},
])

# Reference numbers follow first appearance, including table cells.
order = []
def find_refs(value):
    if isinstance(value, str):
        for number in re.findall(r"\[(\d+)\]", value):
            if int(number) not in order:
                order.append(int(number))
    elif isinstance(value, dict):
        for val in value.values(): find_refs(val)
    elif isinstance(value, list):
        for val in value: find_refs(val)
find_refs(blocks)
mapping = {old: i+1 for i, old in enumerate(order)}
def remap(value):
    if isinstance(value, str):
        return re.sub(r"\[(\d+)\]", lambda m: f"[{mapping[int(m[1])]}]", value)
    if isinstance(value, list): return [remap(v) for v in value]
    if isinstance(value, dict): return {k: remap(v) for k, v in value.items()}
    return value
blocks = remap(blocks)
reference_lookup = {r["number"]: r for r in refs}
references = [dict(reference_lookup[old], original_number=old, number=mapping[old]) for old in order]
(HERE / "references.json").write_text(json.dumps(references, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
chemical_meta=remap(json.loads((HERE/"chemical_references.json").read_text(encoding="utf-8")))
for ref in chemical_meta['references']:
    ref['original_number']=ref['number']
    ref['number']=mapping[ref['number']]
(HERE/'chemical_references_final.json').write_text(json.dumps(chemical_meta,ensure_ascii=False,indent=2)+"\n", encoding="utf-8")

def format_reference(r):
    authors = ", ".join(r["authors"][:3]) + (", et al." if len(r["authors"])>3 else ".")
    vol = r.get("volume", "") + (f"({r['issue']})" if r.get("issue") else "")
    source = f"{r['journal']}, {vol}: {r['pages']}, {r['year']}."
    link = f"DOI: {r['doi']}." if r.get("doi") else f"{r['url']}，访问日期2026-10-23。"
    return f"[{r['number']}] {authors} {r['title']}[J]. {source} {link}"

blocks += [heading("使用声明"), paragraph((ROOT / "docs/use_statement.md").read_text(encoding="utf-8").strip().splitlines()[-1]), heading("参考文献")]
blocks += [{"type": "reference", "text": format_reference(r)} for r in references]
appendix_start = len(blocks)
blocks += [heading("附录A 感官字段及逐维精度"), paragraph("表A-1列出第一问ExtraTrees的全部45维折外误差。A_前缀表示香气，F_前缀为相应风味；其余字段按附件字典保留。所有值均为附件标准化尺度，完整250×45折外预测随电子附录保存。")]
metrics = pd.read_csv(ROOT / "outputs/q1/metrics_by_target.csv")
metrics = metrics.loc[metrics.model.eq("extra_trees")]
blocks.append(table("表A-1 全部感官维度的五折折外指标", ["目标字段", "MAE", "RMSE", "R²", "Spearman"],
                    [[r.target, f"{r.mae:.4f}", f"{r.rmse:.4f}", f"{r.macro_r2:.4f}", f"{r.spearman:.4f}"] for r in metrics.itertuples()],
                    [6.0, 2.5, 2.5, 2.5, 2.5], True))
blocks += [heading("附录B 完整Top 3候选"), paragraph("表B-1沿用绝对Spearman相关排名，不按文献证据后验改选。每项依次给出附件字段名、带符号相关ρ与Bootstrap入选频率π。化学名称、来源、单位、阈值及逐行证据等级见表B-2及电子附件chemical_evidence_final.csv；135个位置涉及60种不同化合物。")]
top = pd.read_csv(ROOT / "outputs/q2/top3_compounds.csv")
rows = []
for target in metrics.target:
    selected = top.loc[top.sensory_target.eq(target)].sort_values("rank")
    assert len(selected)==3
    rows.append([target] + [f"{r.compound}\nρ={r.spearman_rho:.3f}，π={r.bootstrap_top3_frequency:.3f}" for r in selected.itertuples()])
blocks.append(table("表B-1 45项感官的全部Top-3关联候选", ["感官字段", "第1项", "第2项", "第3项"], rows, [3.1, 4.3, 4.3, 4.3], True))
evidence = pd.read_csv(HERE / "chemical_evidence.csv").fillna("")
for col in [c for c in evidence.columns if "ref" in c]:
    evidence[col] = evidence[col].map(lambda x: remap(x) if isinstance(x, str) else x)
for col in evidence.columns:
    evidence[col]=evidence[col].map(lambda x: x.replace('chemical_references.json','chemical_references_final.json') if isinstance(x,str) else x)
evidence.to_csv(HERE / "chemical_evidence_final.csv", index=False, encoding="utf-8-sig", lineterminator="\n")
blocks += [heading("附录C 逐源迁移检查"), paragraph("表C-1采用固定树、中央6004池、retain80、余量优先口径。两个Q95为预留当前模型及新种子旧协议模型的配对预测差分位数；任一大于零均记为失败。改善率已包含所有选中谱，不删除失败的beer 96。完整72点策略网格及十情景化学检查随电子附录保存。")]
transfer = pd.read_csv(ROOT / "outputs/q3_refined/transfer/transfer_summary.csv")
selected = transfer.loc[transfer.arm.eq("fixed") & transfer.pool.eq("contracted_6004") & transfer.bound.eq("retain80") & transfer.ranking.eq("margin")]
blocks.append(table("表C-1 全部11款源的余量优先迁移检查", ["beer", "距离改善", "保留锚幅度", "当前Q95", "旧协议Q95", "审计结果"],
                    [[str(r.source_beer), f"{r.target_improvement:.2%}", f"{r.retained_anchor_gain:.2%}", f"{r.new_audit_q95:+.5f}", f"{r.old_fresh_audit_q95:+.5f}", "通过" if r.fresh_consistency_passed else "未通过"] for r in selected.itertuples()],
                    [1.3, 2.5, 3.2, 3.0, 3.0, 2.5], True))
blocks = add_submission_appendices(blocks, ROOT, HERE, mapping)

# Save a structured editable source with every table included.
manuscript = dict(title=core["title"], abstract=core["abstract"], keywords=core["keywords"], blocks=blocks)
(HERE / "manuscript.json").write_text(json.dumps(manuscript, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")

# Document the source-derived page system before editing its working copy.
with ZipFile(REFERENCE) as z:
    inventory = {name: hashlib.sha256(z.read(name)).hexdigest() for name in z.namelist() if not name.endswith("/")}
template_sha = hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
(HERE / "qa/template_inventory.json").write_text(json.dumps(inventory, indent=2)+"\n", encoding="utf-8")
(HERE / "qa/artifact.md").write_text(f"""# Paper template execution contract
Reference: {REFERENCE}\nSHA-256: {template_sha}\nReference render: qa/template/page-1.png through page-3.png.
The retained source has three portrait A4 sections, 2.00025 cm top/bottom and 2.50119 cm side margins, footer distance 1.74978 cm.
Source roles: centered 14 pt title and contents heading; 12 pt abstract/body; Chinese serif body and Times New Roman Latin text. Existing style and opaque package parts are retained.
Editable: main body placeholder paragraphs and identity table; section header/footer text; styles used for manuscript headings and tables; document properties; settings; image relationships. Replace all placeholder content, retain the three-section abstract/contents/body system, and restart body page numbers at one.
Preserve-only: font table, theme and customXml parts; inventory in template_inventory.json. The original template is never changed.
Source template is sparse and has no completed mathematical/table/figure content. Added components use 12 pt body with fixed 18 pt spacing, 14/12 pt headings, 9.5/10 pt tables with fixed 13 pt spacing, native OMML equations and 300 dpi figures. Table cells wrap and repeat headers. No fixed row heights.
The identity placeholders and source page decorations are removed to satisfy anonymous paper and no-header rules. Chinese font uses the installed Windows SimSun/SimHei and equations use Cambria Math. The source document grid is disabled for content paragraphs because its interaction with the CJK font doubled line spacing; fixed 18 pt body spacing is used instead, with unchanged page geometry.
Title/abstract occupy section one, static linked contents section two, body and appendices section three. Contents page caches are computed from the rendered draft and PAGEREF fields are marked for Word refresh. The revision exports through a separate LibreOffice runtime with fonts_macos.conf; the Windows COM route remains available. Render-derived contents caches are reconciled after export.
Slots: source body paragraphs 3/4/5/22 are title/abstract/keywords; 23–35 are contents placeholders and break; 36 onward is body. These slots may expand. All source placeholder and table text is intentionally removed.
Fidelity gates: template hash, A4 geometry, three sections, anonymous headers/footers, no placeholders, intact preserve-only parts, every final page visually checked, formulas and references reconciled.
""", encoding="utf-8")

doc = Document(REFERENCE)
sections = [deepcopy(s._sectPr) for s in doc.sections]
for spec in sections:
    grid=spec.find(qn("w:docGrid"))
    if grid is not None: spec.remove(grid)
body = doc._element.body
for element in list(body): body.remove(element)
body.append(sections[-1])

def style(name, size, bold=False, family="SimSun"):
    if name not in doc.styles:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s = doc.styles[name]
    s.font.name = "Times New Roman"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = RGBColor(0, 0, 0)
    fonts = s.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), family)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    for key in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme"]:
        fonts.attrib.pop(qn("w:"+key), None)
    pf = s.paragraph_format
    pf.line_spacing = 1.25
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.widow_control = True
    snap=OxmlElement("w:snapToGrid");snap.set(qn("w:val"),"0");s.element.get_or_add_pPr().append(snap)
    return s

normal = style("Normal", 12)
normal.paragraph_format.line_spacing = Pt(18)
normal.paragraph_format.first_line_indent = Pt(24)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
style("Title", 14, True).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
for level, size in [(1, 14), (2, 12), (3, 12)]:
    s=style(f"Heading {level}", size, True, "SimHei")
    s.paragraph_format.first_line_indent = Pt(0)
    s.paragraph_format.space_before = Pt(12 if level==1 else 7)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.keep_together = True
    outline=OxmlElement("w:outlineLvl");outline.set(qn("w:val"), str(level-1));s.element.get_or_add_pPr().append(outline)
for name, size in [("Caption",10.5),("TableText",10),("ReferenceText",10.5),("ContentsText",10.5)]:
    s=style(name,size)
    s.paragraph_format.first_line_indent=Pt(0)
style("Equation",11).paragraph_format.first_line_indent=Pt(0)

def add_text(text, sty=None):
    x=doc.add_paragraph(style=sty)
    x.add_run(text)
    return x

def section_break(spec):
    x=doc.add_paragraph()
    x.paragraph_format.space_after=Pt(0)
    x.paragraph_format.line_spacing=1
    x.paragraph_format.first_line_indent=Pt(0)
    x._p.get_or_add_pPr().append(deepcopy(spec))
    return x

def bookmark(x, name, number):
    start=OxmlElement("w:bookmarkStart");start.set(qn("w:id"),str(number));start.set(qn("w:name"),name)
    end=OxmlElement("w:bookmarkEnd");end.set(qn("w:id"),str(number))
    x._p.insert(0,start);x._p.append(end)

def field(x, instruction, display):
    node=OxmlElement("w:fldSimple");node.set(qn("w:instr"),instruction)
    r=OxmlElement("w:r");t=OxmlElement("w:t");t.text=str(display);r.append(t);node.append(r);x._p.append(node)

title=add_text(core["title"],"Title")
title.paragraph_format.space_after=Pt(12)
title.paragraph_format.first_line_indent=Pt(0)
abstract_heading=add_text("摘 要","Heading 2")
abstract_heading.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in core["abstract"]: add_text(text)
keywords=add_text("关键词："+core["keywords"])
keywords.paragraph_format.first_line_indent=Pt(0)
keywords._p.get_or_add_pPr().append(deepcopy(sections[0]))

toc_title=add_text("目 录","Title")
toc_title.paragraph_format.space_after=Pt(15)
headings=[]
for i,b in enumerate(blocks):
    if b["type"]=="heading":
        text=b["text"]
        if b["level"]==1 or (b["level"]==2 and re.match(r"[567]\.\d+ ",text)):
            headings.append((i,b,f"sec_{i}"))
page_map_path=HERE/"qa/page_map.json"
page_map=json.loads(page_map_path.read_text(encoding="utf-8")) if page_map_path.exists() else {}
for i,b,name in headings:
    x=add_text("", "ContentsText")
    x.paragraph_format.line_spacing=Pt(13)
    x.paragraph_format.space_after=Pt(0)
    x.paragraph_format.left_indent=Cm(.45 if b["level"]==2 else 0)
    x.paragraph_format.tab_stops.add_tab_stop(Cm(15.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    link=OxmlElement("w:hyperlink");link.set(qn("w:anchor"),name)
    r=OxmlElement("w:r");t=OxmlElement("w:t");t.text=b["text"];r.append(t);link.append(r);x._p.append(link)
    x.add_run("\t")
    field(x,f"PAGEREF {name} \\h",page_map.get(b["text"],""))
section_break(sections[1])

def add_table(b):
    cap=add_text(b["caption"],"Caption")
    cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next=True
    widths=b.get("widths")
    if not widths:
        n=len(b["headers"])
        widths=([6]+[10/(n-1)]*(n-1)) if n>2 else [4.5,11.5]
    tbl=doc.add_table(rows=1, cols=len(b["headers"]))
    tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit=False
    for col,w in zip(tbl.columns,widths): col.width=Cm(w)
    allrows=[b["headers"]]+b["rows"]
    for n,values in enumerate(allrows):
        row=tbl.rows[0] if n==0 else tbl.add_row()
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement("w:cantSplit");trpr.append(cant)
        if n==0:
            repeat=OxmlElement("w:tblHeader");trpr.append(repeat)
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"EDEFF2" if n==0 else "FFFFFF");pr.append(shade)
            borders=OxmlElement("w:tcBorders")
            for side in ["top","left","bottom","right"]:
                edge=OxmlElement("w:"+side);edge.set(qn("w:val"),"single");edge.set(qn("w:sz"),"4");edge.set(qn("w:color"),"D9D9D9");borders.append(edge)
            pr.append(borders)
            margins=OxmlElement("w:tcMar")
            for side,amount in [("top",65),("bottom",65),("left",80),("right",80)]:
                v=OxmlElement("w:"+side);v.set(qn("w:w"),str(amount));v.set(qn("w:type"),"dxa");margins.append(v)
            pr.append(margins)
            x=cell.paragraphs[0];x.style=doc.styles["TableText"]
            x.paragraph_format.line_spacing=Pt(13);x.paragraph_format.space_after=Pt(0)
            x.paragraph_format.keep_with_next=n==0 or (len(b["rows"])<=4 and n<len(b["rows"]))
            val=str(value)
            x.alignment=WD_ALIGN_PARAGRAPH.CENTER if (j>0 and len(val)<20) or n==0 else WD_ALIGN_PARAGRAPH.LEFT
            r=x.add_run(val);r.font.size=Pt(9.5 if b.get("compact") else 10)
            r.bold=n==0
    # Keep the gap after the table small and explicit.
    x=doc.add_paragraph();x.paragraph_format.space_after=Pt(3);x.paragraph_format.line_spacing=.3
    x.paragraph_format.first_line_indent=Pt(0)

bookmark_records=[]
for index,b in enumerate(blocks):
    kind=b["type"]
    if kind=="heading":
        x=add_text(b["text"],f"Heading {b['level']}")
        if b['text'].startswith('附录'): x.paragraph_format.page_break_before=True
        bookmark(x,f"sec_{index}",1000+index)
        bookmark_records.append({"text":b["text"],"name":f"sec_{index}","level":b["level"]})
    elif kind in ["paragraph","reference"]:
        x=add_text(b["text"],"ReferenceText" if kind=="reference" else None)
        if kind=="reference":
            x.paragraph_format.left_indent=Cm(.65);x.paragraph_format.first_line_indent=Cm(-.65)
    elif kind=="equation":
        x=doc.add_paragraph(style="Equation")
        x.paragraph_format.keep_together=True
        x.paragraph_format.space_before=Pt(3);x.paragraph_format.space_after=Pt(5)
        x.paragraph_format.tab_stops.add_tab_stop(Cm(7.8),WD_TAB_ALIGNMENT.CENTER)
        x.paragraph_format.tab_stops.add_tab_stop(Cm(15.8),WD_TAB_ALIGNMENT.RIGHT)
        x.add_run("\t");x._p.append(to_omml(b["latex"]));x.add_run("\t（"+b["label"]+"）")
    elif kind=="code":
        x=doc.add_paragraph()
        x.paragraph_format.first_line_indent=Pt(0)
        x.paragraph_format.line_spacing=Pt(12)
        x.paragraph_format.keep_together=True
        r=x.add_run(b["text"])
        r.font.name="Consolas";r.font.size=Pt(9)
    elif kind=="table": add_table(b)
    elif kind=="figure":
        fig=FIGURES[b["key"]]
        x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER
        x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.keep_with_next=True
        x.paragraph_format.line_spacing=1
        x.paragraph_format.space_before=Pt(0);x.paragraph_format.space_after=Pt(3)
        x.add_run().add_picture(str(HERE/fig["relativepath"]),width=Cm(15 if b['key']=='figure5_profiles' else 16))
        caption=fig["caption"].replace("预先列定的", "选定的")
        if b['key']=='figure4_generation':
            caption='图4 三种风格的生成能量差异。柱为外层六折均值，误差线为折间标准差，非置信区间；每折先平均两个模拟种子。'
        if b['key']=='figure5_profiles':
            caption='图5 三种风格的12维代表感官画像。点为预测均值，线为每风格300条生成谱的5%—95%范围，并非预测置信区间。Stale hops用非Pils单目标树，其余维度用226维多输出代理；均为标准化尺度。'
        x=add_text(caption,"Caption");x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        x.paragraph_format.keep_together=True
    else: raise ValueError(kind)

# Clear source header/footer placeholders and use centered body-only numbering.
for i,section in enumerate(doc.sections):
    section.header.is_linked_to_previous=False
    section.footer.is_linked_to_previous=False
    for part in [section.header,section.footer]:
        for node in list(part._element): part._element.remove(node)
        part.add_paragraph()
    if i==2:
        x=section.footer.paragraphs[0];x.alignment=WD_ALIGN_PARAGRAPH.CENTER
        x.paragraph_format.first_line_indent=Pt(0)
        field(x,"PAGE",1)
    for hdr in section._sectPr.findall(qn("w:headerReference")):
        section._sectPr.remove(hdr)
    pn=section._sectPr.find(qn("w:pgNumType"))
    if i==2:
        if pn is None: pn=OxmlElement("w:pgNumType");section._sectPr.append(pn)
        pn.set(qn("w:start"),"1")
settings=doc.settings.element
update=settings.find(qn("w:updateFields"))
if update is None: update=OxmlElement("w:updateFields");settings.append(update)
update.set(qn("w:val"),"true")
mathpr=settings.find(qn("m:mathPr"))
if mathpr is None: mathpr=OxmlElement("m:mathPr");settings.append(mathpr)
font=mathpr.find(qn("m:mathFont"))
if font is None: font=OxmlElement("m:mathFont");mathpr.append(font)
font.set(qn("m:val"),"Cambria Math")
doc.core_properties.title=core["title"]
doc.core_properties.author=""
doc.core_properties.last_modified_by=""
doc.core_properties.comments=""
doc.core_properties.subject="啤酒风味建模"
doc.core_properties.keywords=core["keywords"]
doc.save(FINAL)

# python-docx normalizes relationship XML even for untouched custom parts.
# Restore explicitly preserve-only package members from the retained template.
preserve=[name for name in inventory if name.startswith('customXml/') or name.startswith('word/theme/') or name=='word/fontTable.xml']
temporary=FINAL.with_suffix('.tmp.docx')
with ZipFile(REFERENCE) as original, ZipFile(FINAL) as built, ZipFile(temporary,'w') as out:
    for info in built.infolist():
        data=original.read(info.filename) if info.filename in preserve else built.read(info.filename)
        out.writestr(info,data)
temporary.replace(FINAL)

assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest()==template_sha
with ZipFile(FINAL) as z:
    for name in preserve:
        assert hashlib.sha256(z.read(name)).hexdigest()==inventory[name],name
    xml=z.read('word/document.xml').decode()
    assert len(re.findall(r'<w:sectPr(?:\s|>)',xml))==3
    assert xml.count('<m:oMath>')==sum(b['type']=='equation' for b in blocks)
    for old in ['论文题目','摘要正文','这里开始论文正文','MCB260XXXX','键入章级']:
        assert old not in xml,old
(HERE/"qa/bookmarks.json").write_text(json.dumps(bookmark_records,ensure_ascii=False,indent=2)+"\n", encoding="utf-8")
print(f"Created {FINAL.name}: {len(blocks)} blocks, {sum(b['type']=='equation' for b in blocks)} native equations, {len(references)} references.")

